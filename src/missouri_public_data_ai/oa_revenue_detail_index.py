"""Build and query selected OA General Revenue Detail workbooks."""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
import time
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import requests
from openpyxl import load_workbook

from missouri_public_data_ai.oa_budget_index import build_oa_budget_index


PROJECT_ROOT = Path(__file__).resolve().parents[2]
RAW_DIR = PROJECT_ROOT / "data" / "raw_public" / "oa_revenue_detail"
INDEX_PATH = RAW_DIR / "oa_revenue_detail_index.json"
REPORT_PATH = PROJECT_ROOT / "reports" / "oa_revenue_detail_index_report.json"
SOURCE_PAGE = "https://budplan.oa.mo.gov/revenue-information"
SOURCE_LABEL = "Office of Administration Budget and Planning Revenue Information"

MONTHS = {
    "january": 1,
    "february": 2,
    "march": 3,
    "april": 4,
    "may": 5,
    "june": 6,
    "july": 7,
    "august": 8,
    "september": 9,
    "october": 10,
    "november": 11,
    "december": 12,
}
MONTH_NAMES = {number: name.title() for name, number in MONTHS.items()}

VALUE_FIELDS = {
    "month_current": 1,
    "month_prior": 3,
    "month_percent_change": 5,
    "fiscal_ytd_current": 8,
    "fiscal_ytd_prior": 10,
    "fiscal_ytd_percent_change": 12,
    "three_month_current": 15,
    "three_month_prior": 17,
    "three_month_percent_change": 19,
}

METRIC_ALIASES = [
    (
        "Total Collections Net of Refunds",
        [
            r"\btotal\s+collections\s+net\s+of\s+refunds\b",
            r"\bnet\s+of\s+refunds\b",
            r"\bnet\s+general\s+revenue\b",
            r"\bnet\s+revenue\b",
        ],
    ),
    ("Total Refunds", [r"\btotal\s+refunds\b"]),
    ("Refund Expenditures", [r"\brefund\s+expenditures?\b"]),
    ("Debt Offset Escrow", [r"\bdebt\s+offset\b"]),
    ("Total Collections", [r"\btotal\s+collections\b", r"\bgross\s+collections\b"]),
    ("Sales and Use Tax", [r"\bsales\s+and\s+use\s+tax\b", r"\bsales\s+tax\b"]),
    ("Income Tax - Individual", [r"\bindividual\s+income\s+tax\b", r"\bincome\s+tax\s+-?\s+individual\b"]),
    ("Pass Through Entity Tax", [r"\bpass[-\s]+through\s+entity\s+tax\b"]),
    ("Corporate Income and Franchise Tax", [r"\bcorporate\s+income\b", r"\bfranchise\s+tax\b"]),
    ("County Foreign Insurance", [r"\bcounty\s+foreign\s+insurance\b"]),
    ("Liquor", [r"\bliquor\b"]),
    ("Beer", [r"\bbeer\b"]),
    ("Inheritance/Estate", [r"\binheritance\b", r"\bestate\s+tax\b"]),
    ("All Other Taxes", [r"\ball\s+other\s+taxes\b"]),
    ("Interest", [r"\binterest\b"]),
    ("Licenses, Fees, & Permits", [r"\blicenses?\b.*\bfees?\b", r"\bfees?\b.*\bpermits?\b"]),
    ("Sales, Services, Leases, & Rentals", [r"\bsales,\s*services\b", r"\bleases?\b.*\brentals?\b"]),
    ("Refunds", [r"\brefunds\b"]),
    ("Interagency Billings/Inventory", [r"\binteragency\b", r"\binventory\b"]),
    ("All Other Receipts", [r"\ball\s+other\s+receipts\b"]),
    ("Other", [r"\bother\b"]),
]


def utc_now() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, indent=2, sort_keys=True), encoding="utf-8")


def sha256_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def clean_text(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def normalize_key(value: Any) -> str:
    cleaned = re.sub(r"[^a-z0-9]+", "_", clean_text(value).lower())
    return re.sub(r"_+", "_", cleaned).strip("_")


def value_to_float(value: Any) -> float | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = clean_text(value)
    if not text or text.startswith("#"):
        return None
    text = text.replace("$", "").replace(",", "").replace("%", "")
    try:
        return float(text)
    except ValueError:
        return None


def serializable_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.date().isoformat()
    return value


def month_label_from_text(value: str) -> str | None:
    match = re.search(
        r"\b(January|February|March|April|May|June|July|August|September|October|November|December)\s+\d{1,2},?\s*(20\d{2})\b",
        value,
        flags=re.I,
    )
    if match:
        return f"{match.group(1).title()} {match.group(2)}"
    match = re.search(
        r"\b(January|February|March|April|May|June|July|August|September|October|November|December)\s+(20\d{2})\b",
        value,
        flags=re.I,
    )
    if match:
        return f"{match.group(1).title()} {match.group(2)}"
    return None


def period_key(month_label: str) -> str:
    month, year = month_label.split()
    return f"{year}-{MONTHS[month.lower()]:02d}"


def fiscal_year_for_month(month_label: str) -> int:
    month, year_text = month_label.split()
    year = int(year_text)
    return year + 1 if MONTHS[month.lower()] >= 7 else year


def month_sort_key(month_label: str | None) -> str:
    return period_key(month_label) if month_label else "0000-00"


def safe_filename(source: dict[str, Any], content: bytes) -> str:
    label = source.get("month_label") or source.get("label") or "revenue_detail"
    suffix = Path(str(source.get("url", ""))).suffix.lower() or ".xlsx"
    digest = sha256_bytes(content)[:10]
    return f"{month_sort_key(label)}_{normalize_key(label)}_{digest}{suffix}"


def requested_month_label(question: str, records: list[dict[str, Any]]) -> str | None:
    lowered = question.lower()
    years = [int(year) for year in re.findall(r"\b(20\d{2})\b", question)]
    for month_name in MONTHS:
        if re.search(rf"\b{month_name}\b", lowered):
            if years:
                return f"{month_name.title()} {years[-1]}"
            matching = [row["month_label"] for row in records if row.get("month_label", "").lower().startswith(month_name)]
            return sorted(matching, key=period_key)[-1] if matching else None
    fy_match = re.search(r"\bFY\s*20?(\d{2,4})\b", question, flags=re.I)
    fiscal_year = None
    if fy_match:
        raw = fy_match.group(1)
        fiscal_year = int(raw) if len(raw) == 4 else int(f"20{raw}")
    elif "fiscal year" in lowered and years:
        fiscal_year = years[-1]
    if fiscal_year:
        matching = [row["month_label"] for row in records if row.get("fiscal_year") == fiscal_year]
        return sorted(set(matching), key=period_key)[-1] if matching else None
    if any(term in lowered for term in ["latest", "newest", "most recent", "current"]):
        return sorted({row["month_label"] for row in records if row.get("month_label")}, key=period_key)[-1]
    return None


def requested_metric_key(question: str) -> str | None:
    lowered = question.lower()
    if (
        "general revenue" in lowered
        and not any(term in lowered for term in ["available", "link", "where", "source"])
        and any(term in lowered for term in ["what", "how much", "amount", "collections", "collected", "was", "were"])
    ):
        return normalize_key("Total Collections Net of Refunds")
    for label, patterns in METRIC_ALIASES:
        if any(re.search(pattern, lowered) for pattern in patterns):
            return normalize_key(label)
    return None


def requested_value_field(question: str) -> str:
    lowered = question.lower()
    percent = any(term in lowered for term in ["percent", "percentage", "increase", "decrease", "change"])
    if any(term in lowered for term in ["year-to-date", "year to date", "ytd", "fiscal year to date", "fytd"]):
        return "fiscal_ytd_percent_change" if percent else "fiscal_ytd_current"
    if any(term in lowered for term in ["last 3 months", "last three months", "three month", "trailing"]):
        return "three_month_percent_change" if percent else "three_month_current"
    if percent:
        return "month_percent_change"
    return "month_current"


def format_money(value: float | None) -> str:
    return "not reported" if value is None else f"${value:,.2f}"


def format_percent(value: float | None) -> str:
    return "not reported" if value is None else f"{value:,.1f}%"


def field_label(field: str) -> str:
    return {
        "month_current": "monthly amount",
        "month_prior": "prior-year monthly amount",
        "month_percent_change": "monthly percent change",
        "fiscal_ytd_current": "fiscal year-to-date amount",
        "fiscal_ytd_prior": "prior-year fiscal year-to-date amount",
        "fiscal_ytd_percent_change": "fiscal year-to-date percent change",
        "three_month_current": "last 3 months amount",
        "three_month_prior": "prior-year last 3 months amount",
        "three_month_percent_change": "last 3 months percent change",
    }[field]


def discover_detail_sources(metadata: dict[str, Any]) -> list[dict[str, Any]]:
    rows = [
        row
        for row in metadata.get("records", [])
        if row.get("document_type") == "revenue_detail" and row.get("file_type") == "xlsx" and row.get("month_label")
    ]
    return sorted(rows, key=lambda item: month_sort_key(item.get("month_label")))


def parse_workbook(content: bytes, source: dict[str, Any], local_path: Path) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    worksheet = workbook["OTHER"] if "OTHER" in workbook.sheetnames else workbook.worksheets[0]
    month_label = source.get("month_label")
    date_prepared = None
    records: list[dict[str, Any]] = []
    for row_number, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
        label = clean_text(row[0] if row else "")
        if not label:
            continue
        if "MONTH ENDED" in label.upper():
            month_label = month_label_from_text(label) or month_label
            continue
        if label.upper().startswith("DATE PREPARED"):
            date_prepared = serializable_value(row[1] if len(row) > 1 else None)
            continue
        values = {field: value_to_float(row[index] if index < len(row) else None) for field, index in VALUE_FIELDS.items()}
        if all(value is None for value in values.values()):
            continue
        if any(
            label.upper().startswith(prefix)
            for prefix in [
                "PREPARED BY",
                "MISSOURI DEPARTMENT",
                "ADMINISTRATION DIVISION",
                "FINANCIAL SERVICES",
                "MONTHLY GENERAL",
                "SOURCE:",
            ]
        ):
            continue
        records.append(
            {
                "month_label": month_label,
                "period_key": period_key(month_label) if month_label else None,
                "fiscal_year": fiscal_year_for_month(month_label) if month_label else None,
                "metric_label": label,
                "metric_key": normalize_key(label),
                "sheet": worksheet.title,
                "row_number": row_number,
                "source_url": source["url"],
                "source_label": source.get("label"),
                "local_file": str(local_path.relative_to(PROJECT_ROOT)),
                **values,
            }
        )
    file_record = {
        "month_label": month_label,
        "period_key": period_key(month_label) if month_label else None,
        "fiscal_year": fiscal_year_for_month(month_label) if month_label else None,
        "source_label": source.get("label"),
        "source_url": source["url"],
        "source_page": source.get("source_page") or SOURCE_PAGE,
        "local_file": str(local_path.relative_to(PROJECT_ROOT)),
        "sheet": worksheet.title,
        "bytes": len(content),
        "sha256": sha256_bytes(content),
        "date_prepared": date_prepared,
        "metric_count": len(records),
    }
    return file_record, records


def build_oa_revenue_detail_index(force: bool = False) -> dict[str, Any]:
    if INDEX_PATH.exists() and not force:
        return json.loads(INDEX_PATH.read_text(encoding="utf-8"))

    metadata = build_oa_budget_index(force=force)
    detail_sources = discover_detail_sources(metadata)
    session = requests.Session()
    session.headers.update(
        {
            "User-Agent": "Mozilla/5.0 (compatible; MissouriPublicDataAssistant/1.0)",
            "Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,*/*;q=0.8",
        }
    )
    RAW_DIR.mkdir(parents=True, exist_ok=True)
    start = time.perf_counter()
    source_files: list[dict[str, Any]] = []
    records: list[dict[str, Any]] = []
    errors: list[dict[str, str]] = []
    for source in detail_sources:
        try:
            response = session.get(source["url"], timeout=60)
            response.raise_for_status()
            content = response.content
            local_path = RAW_DIR / safe_filename(source, content)
            local_path.write_bytes(content)
            file_record, file_records = parse_workbook(content, source, local_path)
            source_files.append(file_record)
            records.extend(file_records)
        except Exception as exc:  # pragma: no cover - surfaced in build report
            errors.append({"label": source.get("label", ""), "url": source.get("url", ""), "error": str(exc)})

    records = sorted(records, key=lambda item: (item.get("period_key") or "", item.get("row_number") or 0))
    source_files = sorted(source_files, key=lambda item: item.get("period_key") or "")
    metric_counts = Counter(row["metric_label"] for row in records)
    fiscal_years = sorted({row["fiscal_year"] for row in records if row.get("fiscal_year")})
    period_keys = [item["period_key"] for item in source_files if item.get("period_key")]
    payload = {
        "generated_at_utc": utc_now(),
        "elapsed_seconds": round(time.perf_counter() - start, 3),
        "source": SOURCE_LABEL,
        "source_page": SOURCE_PAGE,
        "index_path": str(INDEX_PATH.relative_to(PROJECT_ROOT)),
        "workbook_count": len(source_files),
        "record_count": len(records),
        "downloaded_bytes": sum(item.get("bytes", 0) for item in source_files),
        "downloaded_mb": round(sum(item.get("bytes", 0) for item in source_files) / 1_000_000, 3),
        "min_period": min(period_keys) if period_keys else None,
        "max_period": max(period_keys) if period_keys else None,
        "fiscal_years": fiscal_years,
        "metric_counts": dict(sorted(metric_counts.items())),
        "errors": errors,
        "notes": [
            "This index parses selected monthly FY 2026 OA General Revenue Detail Excel workbooks linked from the official Revenue Information page.",
            "It stores aggregate revenue/refund line items only, not taxpayer records.",
            "Older final fiscal-year detail PDFs and broader budget documents still need separate parsers.",
        ],
        "source_files": source_files,
        "records": records,
    }
    write_json(INDEX_PATH, payload)
    write_json(REPORT_PATH, {key: value for key, value in payload.items() if key != "records"})
    return payload


class OaRevenueDetailIndex:
    def __init__(self, path: Path = INDEX_PATH) -> None:
        self.path = path
        self._payload: dict[str, Any] | None = None

    def available(self) -> bool:
        return self.path.exists() and self.path.stat().st_size > 0

    def payload(self) -> dict[str, Any]:
        if self._payload is None:
            self._payload = json.loads(self.path.read_text(encoding="utf-8")) if self.available() else {}
        return self._payload

    def records(self) -> list[dict[str, Any]]:
        return list(self.payload().get("records", []))

    def source_files(self) -> list[dict[str, Any]]:
        return list(self.payload().get("source_files", []))

    def citation(self, matched_rows: int = 0, source_file: dict[str, Any] | None = None) -> list[dict[str, Any]]:
        payload = self.payload()
        file_rows = [source_file] if source_file else self.source_files()[:3]
        source_files = [
            {
                "category": "oa_revenue_detail",
                "category_label": "OA General Revenue Detail workbook",
                "file_name": row.get("source_url"),
                "source_url": row.get("source_url"),
                "row_count": row.get("metric_count"),
                "bytes": row.get("bytes"),
                "sha256": row.get("sha256"),
            }
            for row in file_rows
        ]
        return [
            {
                "dataset": payload.get("source", SOURCE_LABEL),
                "category": "Budget and revenue",
                "kind": "OA general revenue detail workbook",
                "lookup_table": "oa_revenue_detail_index",
                "year": source_file.get("fiscal_year") if source_file else None,
                "year_range": (
                    f"{payload.get('min_period')}-{payload.get('max_period')}"
                    if payload.get("min_period") and payload.get("max_period")
                    else None
                ),
                "source_files": source_files,
                "source_file_count": payload.get("workbook_count"),
                "source_rows": payload.get("record_count"),
                "matched_rows": matched_rows,
            }
        ]

    def unavailable_answer(self, question: str) -> dict[str, Any]:
        return {
            "question": question,
            "answer": (
                "The OA General Revenue Detail workbook index has not been built yet. Run "
                "`python scripts/build_oa_revenue_detail_index.py --force` to parse the official monthly Excel files."
            ),
            "retrieved_context_id": "oa_revenue_detail_index:missing",
            "retrieved_source": "oa_revenue_detail_lookup_index",
            "retrieval_score": 1.0,
            "used_model": False,
            "model": "deterministic_public_lookup",
            "source_note": "The OA revenue-detail route exists, but the local ignored index is missing.",
            "citations": [],
            "source_rows": [],
            "source_url": SOURCE_PAGE,
        }

    def summary_answer(self, question: str) -> dict[str, Any]:
        payload = self.payload()
        return {
            "question": question,
            "answer": (
                "The OA General Revenue Detail exact lookup layer parses official monthly Excel workbooks from the "
                f"Revenue Information page. It currently covers {payload.get('workbook_count', 0):,} workbook(s), "
                f"{payload.get('record_count', 0):,} aggregate revenue/refund line item(s), periods "
                f"{payload.get('min_period')}-{payload.get('max_period')}, and fiscal year(s) "
                f"{', '.join(str(year) for year in payload.get('fiscal_years', []))}."
            ),
            "retrieved_context_id": "oa_revenue_detail_index:summary",
            "retrieved_source": "oa_revenue_detail_lookup_index",
            "retrieval_score": 1.0,
            "used_model": False,
            "model": "deterministic_public_lookup",
            "source_note": "Computed from the local OA General Revenue Detail workbook index.",
            "citations": self.citation(matched_rows=payload.get("record_count", 0)),
            "source_rows": [],
            "source_url": SOURCE_PAGE,
        }

    def find_source_file(self, month_label: str) -> dict[str, Any] | None:
        for row in self.source_files():
            if row.get("month_label") == month_label:
                return row
        return None

    def rows_for_month(self, month_label: str) -> list[dict[str, Any]]:
        return [row for row in self.records() if row.get("month_label") == month_label]

    def file_answer(self, question: str, month_label: str) -> dict[str, Any]:
        rows = self.rows_for_month(month_label)
        source_file = self.find_source_file(month_label)
        if not rows or not source_file:
            return self.missing_answer(question)
        examples = ", ".join(row["metric_label"] for row in rows[:5])
        return {
            "question": question,
            "answer": (
                f"The {month_label} OA General Revenue Detail workbook is parsed with {len(rows)} aggregate line item(s). "
                f"Examples include {examples}. Official workbook: {source_file['source_url']}."
            ),
            "retrieved_context_id": f"oa_revenue_detail_index:file:{source_file.get('period_key')}",
            "retrieved_source": "oa_revenue_detail_lookup_index",
            "retrieval_score": 1.0,
            "used_model": False,
            "model": "deterministic_public_lookup",
            "source_note": "Computed from the local OA General Revenue Detail workbook index.",
            "citations": self.citation(matched_rows=len(rows), source_file=source_file),
            "source_rows": [],
            "source_url": source_file["source_url"],
        }

    def value_answer(self, question: str, record: dict[str, Any], field: str) -> dict[str, Any]:
        month_label = record["month_label"]
        source_file = self.find_source_file(month_label) or {}
        value = record.get(field)
        is_percent = field.endswith("percent_change")
        rendered_value = format_percent(value) if is_percent else format_money(value)
        label = field_label(field)
        metric = record["metric_label"]
        metric_phrase = metric
        if normalize_key(metric) == normalize_key("Total Collections Net of Refunds"):
            metric_phrase = f"net general revenue collections ({metric})"
        if is_percent:
            current_field = field.replace("percent_change", "current")
            prior_field = field.replace("percent_change", "prior")
            comparison = (
                f" The same row shows {format_money(record.get(current_field))} current period versus "
                f"{format_money(record.get(prior_field))} prior period."
            )
        else:
            comparison = ""
        return {
            "question": question,
            "answer": (
                f"For {month_label}, the official OA General Revenue Detail workbook reports {metric_phrase} "
                f"{label} as {rendered_value}.{comparison}"
            ),
            "retrieved_context_id": f"oa_revenue_detail_index:{record.get('period_key')}:{record.get('metric_key')}:{field}",
            "retrieved_source": "oa_revenue_detail_lookup_index",
            "retrieval_score": 1.0,
            "used_model": False,
            "model": "deterministic_public_lookup",
            "source_note": "Computed from a monthly OA General Revenue Detail workbook linked by Budget and Planning.",
            "citations": self.citation(matched_rows=1, source_file=source_file),
            "source_rows": [
                {
                    "source_file": record["source_url"],
                    "source_row_number": record["row_number"],
                    "values": {
                        "month": month_label,
                        "metric": metric,
                        "field": field,
                        "value": round(value, 4) if isinstance(value, float) else value,
                        "monthly_amount": record.get("month_current"),
                        "fiscal_ytd_amount": record.get("fiscal_ytd_current"),
                        "monthly_percent_change": record.get("month_percent_change"),
                    },
                }
            ],
            "source_url": record["source_url"],
        }

    def missing_answer(self, question: str) -> dict[str, Any]:
        payload = self.payload()
        return {
            "question": question,
            "answer": (
                "I have parsed selected OA General Revenue Detail workbooks, but this question did not match an indexed "
                "month/year and revenue line item. Try `What were net general revenue collections in January 2026?`, "
                "`How much Sales and Use Tax did Missouri collect in April 2026?`, or "
                "`What was the FY 2026 year-to-date total collections net of refunds?`"
            ),
            "retrieved_context_id": "oa_revenue_detail_index:no_match",
            "retrieved_source": "oa_revenue_detail_lookup_index",
            "retrieval_score": 1.0,
            "used_model": False,
            "model": "deterministic_public_lookup",
            "source_note": (
                f"Indexed periods: {payload.get('min_period')}-{payload.get('max_period')} from official OA revenue-detail workbooks."
            ),
            "citations": self.citation(),
            "source_rows": [],
            "source_url": SOURCE_PAGE,
        }

    def answer(self, question: str) -> dict[str, Any] | None:
        if not self.available():
            return self.unavailable_answer(question)
        lowered = question.lower()
        records = self.records()
        if any(term in lowered for term in ["indexed", "connected", "coverage", "data"]) and not requested_metric_key(question):
            return self.summary_answer(question)
        month_label = requested_month_label(question, records)
        if not month_label:
            return self.missing_answer(question)
        metric_key = requested_metric_key(question)
        if not metric_key:
            return self.file_answer(question, month_label)
        rows = [row for row in self.rows_for_month(month_label) if row.get("metric_key") == metric_key]
        if not rows:
            return self.missing_answer(question)
        return self.value_answer(question, rows[0], requested_value_field(question))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--force", action="store_true")
    args = parser.parse_args()
    payload = build_oa_revenue_detail_index(force=args.force)
    print(json.dumps({key: value for key, value in payload.items() if key != "records"}, indent=2, sort_keys=True))


if __name__ == "__main__":
    main()
